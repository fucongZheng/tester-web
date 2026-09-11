"""测试用例 Excel/CSV 解析与归一化（纯函数，不依赖 FastAPI）"""
import csv
import io
import re

from openpyxl import Workbook, load_workbook

MAX_ROWS = 1000        # 数据行上限
MAX_CELL = 10000       # 单元格字符上限
MAX_HEADER_SCAN = 5    # 表头探测行数

# 字段 -> 表头别名（经 normalize_header 归一化后精确匹配）；title 是锚点字段
ALIASES = {
    "title": ["标题", "用例标题", "用例名称", "用例", "测试内容"],
    "steps": ["步骤", "测试步骤", "操作步骤", "执行步骤", "用例步骤", "测试过程"],
    "expected": ["预期", "预期结果", "期望结果", "预期输出"],
    "precondition": ["前置条件", "预置条件", "前置"],
    "module": ["模块", "所属模块", "模块名称", "功能模块"],
    "priority": ["优先级", "级别", "优先级别", "用例优先级"],
    "case_type": ["类型", "用例类型", "测试类型", "用例分类"],
    "remark": ["备注", "说明"],
    "origin_no": ["编号", "用例id", "用例编号", "id", "caseid"],  # 只保留进备注，不动系统 case_no
}
FIELD_PRIORITY = list(ALIASES.keys())

CASE_TYPE_ENUM = ["功能", "接口", "UI", "性能", "回归"]
PRIORITY_ALIAS = {"最高": "P0", "高": "P1", "中": "P2", "低": "P3", "最低": "P3"}
SKIP_TITLES = {"-", "无", "/"}

# 全角 ASCII -> 半角
_FW = {chr(0xFF01 + i): chr(0x21 + i) for i in range(94)}


class ParseError(Exception):
    """消息直接作为 400 detail 展示给用户。"""


def normalize_header(v) -> str:
    s = str(v or "").strip()
    s = "".join(_FW.get(ch, ch) for ch in s)
    s = re.sub(r"\s+", "", s)  # 去掉内嵌空白，如「用例 ID」
    return s.lower()


def cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()[:MAX_CELL]


def read_workbook(raw: bytes, ext: str) -> dict:
    """返回 {sheet名: [[单元格str], ...]}；每 sheet 最多读 MAX_ROWS+MAX_HEADER_SCAN 行。"""
    ext = (ext or "").lower()
    if ext == ".xls":
        raise ParseError("暂不支持 .xls，请在 Excel 中另存为 .xlsx 或 .csv 后导入")
    if ext == ".csv":
        text = None
        for enc in ("utf-8-sig", "gb18030"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ParseError("CSV 编码无法识别，请用 UTF-8 重新保存或转为 .xlsx")
        rows = [[cell_to_str(c) for c in r] for r in csv.reader(io.StringIO(text))]
        return {"CSV": rows[: MAX_ROWS + MAX_HEADER_SCAN]}
    if ext == ".xlsx":
        try:
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        except Exception:
            raise ParseError("文件解析失败，请确认为有效的 .xlsx 文件")
        sheets = {}
        try:
            for ws in wb.worksheets:
                rows = []
                for r in ws.iter_rows(values_only=True):
                    rows.append([cell_to_str(c) for c in r])
                    if len(rows) >= MAX_ROWS + MAX_HEADER_SCAN:
                        break
                sheets[ws.title] = rows
        finally:
            wb.close()
        return sheets
    raise ParseError("仅支持 .xlsx / .csv 文件")


def _alias_hits(row) -> int:
    seen = set()
    for cell in row:
        h = normalize_header(cell)
        for field, aliases in ALIASES.items():
            if h in aliases:
                seen.add(field)
                break
    return len(seen)


def detect_header_row(rows) -> int:
    """前 5 行中命中别名词族最多的行（1-based），并列取最早。"""
    best_idx, best_hits = 0, 0
    for i, row in enumerate(rows[:MAX_HEADER_SCAN], 1):
        hits = _alias_hits(row)
        if hits > best_hits:
            best_idx, best_hits = i, hits
    if not best_hits:
        raise ParseError("未识别到表头行（前 5 行内未找到 标题/步骤/预期 等表头），请检查文件格式")
    return best_idx


def auto_map(headers) -> dict:
    """表头 -> {字段: 列索引|None}；按字段优先级占列，一列只归一个字段。"""
    norm = [normalize_header(h) for h in headers]
    mapping = {f: None for f in FIELD_PRIORITY}
    used = set()
    for field in FIELD_PRIORITY:
        aliases = ALIASES[field]
        for idx, h in enumerate(norm):
            if h and h in aliases and idx not in used:
                mapping[field] = idx
                used.add(idx)
                break
    return mapping


def score_sheet(rows) -> int:
    """title 命中=3，steps/expected=2，其余=1；title 未命中即非候选 sheet（0 分）。"""
    try:
        hr = detect_header_row(rows)
    except ParseError:
        return 0
    m = auto_map(rows[hr - 1])
    if m["title"] is None:
        return 0
    score = 3
    score += 2 if m["steps"] is not None else 0
    score += 2 if m["expected"] is not None else 0
    score += sum(1 for f in ("precondition", "module", "priority", "case_type", "remark", "origin_no")
                 if m[f] is not None)
    return score


def norm_priority(v):
    """返回 (P0~P3, 警告文案)。"""
    s = str(v or "").strip()
    if not s:
        return "P2", ""
    compact = re.sub(r"\s+", "", s).upper()
    m = re.fullmatch(r"P([0-3])", compact)
    if m:
        return compact, ""
    if compact in {"0", "1", "2", "3"}:
        return f"P{compact}", ""
    for alias, p in PRIORITY_ALIAS.items():
        if compact == alias.upper() or s == alias:
            return p, f"优先级“{s}”已映射为 {p}"
    return "P2", f"优先级“{s}”无法识别，已按 P2 导入"


def norm_case_type(v):
    """返回 (枚举值, 警告文案, 保留进备注的原值)。"""
    s = str(v or "").strip()
    if not s:
        return "功能", "", ""
    low = s.lower()
    if "接口" in s:
        return "接口", "", ""
    if "ui" in low or "界面" in s:
        return "UI", "", ""
    if "性能" in s:
        return "性能", "", ""
    if "回归" in s:
        return "回归", "", ""
    if "功能" in s:  # 覆盖 功能/功能性
        return "功能", "", ""
    return "功能", "", s


def norm_steps(v) -> str:
    s = str(v or "").strip()
    if not s:
        return ""
    if "\n" in s:  # 已是换行分隔：仅去空行，绝不重新编号
        return "\n".join(ln.strip() for ln in s.splitlines() if ln.strip())
    if " / " in s:
        parts = [p.strip() for p in s.split(" / ")]
        numbered = sum(1 for p in parts if re.match(r"^\d+[\.、\)]", p))
        if numbered >= 2:
            return "\n".join(p for p in parts if p)
    return s


def parse_sheet(rows, header_row: int, mapping: dict):
    """归一化数据行。返回 (行列表, 统计)。跳过的行保留在结果里并带 _skip/_skip_reason。"""
    headers = rows[header_row - 1] if 0 < header_row <= len(rows) else []
    header_norms = {normalize_header(h) for h in headers}

    def cell(row, field):
        idx = mapping.get(field)
        if idx is None or idx < 0 or idx >= len(row):
            return ""
        return row[idx]

    data = rows[header_row:]
    truncated = len(data) > MAX_ROWS
    if truncated:
        data = data[:MAX_ROWS]

    out, valid, skipped, warned = [], 0, 0, 0
    for i, row in enumerate(data):
        title = cell(row, "title").strip()
        priority, warn = norm_priority(cell(row, "priority"))
        case_type, _, type_raw = norm_case_type(cell(row, "case_type"))
        remark_parts = [cell(row, "remark").strip()]
        origin_no = cell(row, "origin_no").strip()
        if origin_no:
            remark_parts.append(f"原编号：{origin_no}")
        if type_raw:
            remark_parts.append(f"Excel类型：{type_raw}")
        remark = "；".join(p for p in remark_parts if p)

        skip, skip_reason = "", ""
        if not any(c.strip() for c in row):
            skip, skip_reason = True, "空行"
        elif any("模块分割行" in c for c in row[:3]):
            skip, skip_reason = True, "模块分割行"
        elif not title or title in SKIP_TITLES:
            skip, skip_reason = True, "标题为空"
        elif normalize_header(title) in header_norms:
            skip, skip_reason = True, "重复表头"

        title_warn = ""
        if not skip and len(title) > 255:
            title_warn = "标题超 255 字已截断"
        rec = {
            "_row": header_row + 1 + i,
            "title": title[:255],
            "precondition": cell(row, "precondition").strip(),
            "steps": norm_steps(cell(row, "steps")),
            "expected": cell(row, "expected").strip(),
            "module": cell(row, "module").strip(),
            "priority": priority,
            "case_type": case_type,
            "case_type_raw": type_raw,
            "remark": remark,
            "origin_no": origin_no,
            "_skip": skip,
            "_skip_reason": skip_reason,
            "_warn": "；".join(w for w in (warn, title_warn) if w),
        }
        if skip:
            skipped += 1
        else:
            valid += 1
            if rec["_warn"]:
                warned += 1
        out.append(rec)

    stats = {"total_rows": len(data), "valid": valid, "skipped": skipped,
             "warned": warned, "truncated": truncated}
    return out, stats


def build_template() -> bytes:
    """标准导入模板 xlsx。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"
    ws.append(["标题", "前置条件", "步骤", "预期结果", "模块", "优先级", "类型", "备注"])
    ws.append(["示例：Logo上传-验证jpg格式上传成功", "机构管理员已登录",
               "1. 单击Logo上传区\n2. 选择 test.jpg", "上传成功并回显缩略图", "应用管理", "P0", "功能", ""])
    ws.append(["示例：Logo上传-验证gif格式被拦截", "机构管理员已登录",
               "1. 单击Logo上传区\n2. 选择 test.gif", "提示格式不符，未回显", "应用管理", "P1", "功能", ""])
    for col, w in zip("ABCDEFGH", (34, 22, 34, 30, 14, 8, 8, 18)):
        ws.column_dimensions[col].width = w

    info = wb.create_sheet("填写说明")
    for r in [
        ["字段", "规则"],
        ["标题", "必填，最长 255 字符；同一项目同一版本下标题重复会被跳过"],
        ["步骤", "建议每步一行，形如 1. xxx（换行）2. xxx"],
        ["优先级", "P0 / P1 / P2 / P3；也支持 高/中/低（映射为 P1/P2/P3）"],
        ["类型", "功能 / 接口 / UI / 性能 / 回归；其他值（如 异常流、安全性）按“功能”导入并在备注保留原值"],
        ["模块", "可选；导入时可选 指定模块 / 按名称匹配系统模块 / 自动新建缺失模块"],
        ["表头", "支持常见别名（用例标题/测试步骤/预期结果/所属模块 等），列顺序不限"],
    ]:
        info.append(r)
    info.column_dimensions["A"].width = 10
    info.column_dimensions["B"].width = 70

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
