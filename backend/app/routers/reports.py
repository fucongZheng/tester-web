"""测试报告：聚合现有数据填入 Word 模板后导出"""
import io
import json
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from jinja2 import Environment, FileSystemLoader
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import (TestReport, Project, Version, Requirement, TestCase,
                      Execution, Bug, FlowInstance, Module)
from ..deps import get_current_user, require_admin
from ..helpers import row_to_dict, gen_code
from ..flow_stages import stage_label

router = APIRouter(prefix="/api/reports", tags=["测试报告"])

TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates"
DOCX_TEMPLATE = TEMPLATE_DIR / "report_template.docx"
# trim_blocks+lstrip_blocks：让 {% for %} 这类块标签整行消失，
# 否则渲染出的 markdown 表格行之间夹空行，GFM 解析时表格会被截断成散文本
_env = Environment(loader=FileSystemLoader(str(TEMPLATE_DIR)),
                   trim_blocks=True, lstrip_blocks=True)

CLOSED = {"测试验证通过关闭", "不是BUG"}
SEV_ORDER = ["致命", "严重", "一般", "轻微", "建议"]


def _pct(n, d):
    return f"{n / d * 100:.1f}%" if d else "0.0%"


def _count(seq, key):
    m = {}
    for x in seq:
        m[key(x)] = m.get(key(x), 0) + 1
    return m


def _strip_html(html: str) -> str:
    import re
    text = re.sub(r"<[^>]+>", "", html or "")
    return text.replace("&nbsp;", " ").strip()


def _aggregate(db, project_id, version_id, module_ids=None):
    """module_ids 非空时报告只统计所选模块（需求/用例/BUG 按 module_id 过滤），空则整版本。"""
    project = db.query(Project).filter(Project.id == project_id).first()
    version = db.query(Version).filter(Version.id == version_id).first()
    if not project or not version:
        raise HTTPException(status_code=400, detail="项目或版本不存在")

    module_ids = [m for m in (module_ids or []) if m]
    req_q = db.query(Requirement).filter(Requirement.version_id == version_id)
    case_q = db.query(TestCase).filter(TestCase.version_id == version_id)
    bug_q = db.query(Bug).filter(Bug.version_id == version_id)
    if module_ids:
        req_q = req_q.filter(Requirement.module_id.in_(module_ids))
        case_q = case_q.filter(TestCase.module_id.in_(module_ids))
        bug_q = bug_q.filter(Bug.module_id.in_(module_ids))
    reqs = req_q.all()
    cases = case_q.all()
    case_ids = [c.id for c in cases]
    execs = db.query(Execution).filter(Execution.case_id.in_(case_ids)).all() if case_ids else []
    bugs = bug_q.all()
    if module_ids:
        modules = db.query(Module).filter(Module.id.in_(module_ids)).order_by(Module.id).all()
    else:
        modules = db.query(Module).filter(Module.project_id == project_id).all()

    latest = {}
    for e in execs:
        prev = latest.get(e.case_id)
        if not prev or (e.id or 0) > (prev.id or 0):
            latest[e.case_id] = e
    latest_list = list(latest.values())
    exec_result = {"通过": 0, "失败": 0, "阻塞": 0, "跳过": 0, "未执行": 0}
    for e in latest_list:
        exec_result[e.result] = exec_result.get(e.result, 0) + 1
    not_executed = exec_result.get("未执行", 0) + max(0, len(cases) - len(latest_list))
    executed_total = len(cases) - not_executed
    pass_count = exec_result.get("通过", 0)
    fail_count = exec_result.get("失败", 0)
    pass_rate = _pct(pass_count, executed_total)
    exec_ratio = {k: _pct(v, len(cases) or 1) for k, v in exec_result.items()}
    exec_ratio["覆盖率"] = _pct(len(latest_list), len(cases))

    req_status = _count(reqs, lambda r: r.status)
    bug_severity = _count(bugs, lambda b: b.severity)
    bug_status = _count(bugs, lambda b: b.status)
    bug_stage = _count(bugs, lambda b: b.found_stage)
    bug_module = {}
    for b in bugs:
        name = b.module.name if b.module else "未分类"
        bug_module[name] = bug_module.get(name, 0) + 1

    closed = [b for b in bugs if b.status in CLOSED]
    open_bugs = [b for b in bugs if b.status not in CLOSED]

    by_mod_case = {}
    for c in cases:
        name = c.module.name if c.module else "未分类"
        d = by_mod_case.setdefault(name, {"total": 0, "pass": 0, "fail": 0})
        d["total"] += 1
        e = latest.get(c.id)
        if e and e.result == "通过":
            d["pass"] += 1
        elif e and e.result == "失败":
            d["fail"] += 1

    flow = db.query(FlowInstance).filter(FlowInstance.version_id == version_id).order_by(
        FlowInstance.id.desc()).first()
    flow_stage = stage_label(flow.current_stage, flow.current_round) if flow else "未启动"
    flow_round = flow.current_round if flow else 1
    flow_status = flow.status if flow else "未启动"

    testers = sorted({e.executor for e in execs if e.executor} | {b.submitter for b in bugs if b.submitter})
    tester_names = "、".join(testers) if testers else ""

    return {
        "project": project, "version": version, "reqs": reqs, "cases": cases,
        "bugs": bugs, "open_bugs": open_bugs, "closed": closed, "modules": modules,
        "req_total": len(reqs), "req_status": req_status,
        "case_total": len(cases), "executed_total": executed_total,
        "not_executed": not_executed, "pass_count": pass_count, "fail_count": fail_count,
        "pass_rate": pass_rate, "exec_result": exec_result, "exec_ratio": exec_ratio,
        "bug_total": len(bugs), "bug_open": len(open_bugs), "bug_closed": len(closed),
        "bug_severity": bug_severity, "bug_status": bug_status,
        "bug_module": bug_module, "bug_stage": bug_stage, "by_mod_case": by_mod_case,
        "flow": flow, "flow_stage": flow_stage, "flow_round": flow_round, "flow_status": flow_status,
        "tester_names": tester_names,
    }


def _md_content(title, agg, generator, generated_at):
    ctx = {
        "title": title,
        "project_name": agg["project"].name,
        "version_no": agg["version"].version_no,
        "version_name": agg["version"].name,
        "start_date": agg["version"].start_date or "-",
        "launch_date": agg["version"].launch_date or "-",
        "generated_at": generated_at,
        "method": "系统",
        "req_total": agg["req_total"], "req_status": agg["req_status"],
        "case_total": agg["case_total"], "executed_total": agg["executed_total"],
        "not_executed": agg["not_executed"], "pass_rate": agg["pass_rate"],
        "exec_result": agg["exec_result"], "exec_ratio": agg["exec_ratio"],
        "bug_total": agg["bug_total"], "bug_severity": agg["bug_severity"],
        "bug_status": agg["bug_status"], "bug_module": agg["bug_module"],
        "bug_stage": agg["bug_stage"],
        "flow_stage": agg["flow_stage"], "flow_round": agg["flow_round"],
        "modules": agg["modules"],
    }
    return _env.get_template("report.md.j2").render(**ctx)


def _set_cell(cell, text):
    text = "" if text is None else str(text)
    if cell.paragraphs:
        runs = cell.paragraphs[0].runs
        if runs:
            runs[0].text = text
            for r in runs[1:]:
                r.text = ""
            for p in cell.paragraphs[1:]:
                p.clear()
            return
        cell.paragraphs[0].text = text
        return
    cell.text = text


def _fill_table(table, rows, start=1):
    """保留表头，从 start 行起用 rows 覆盖/追加。rows 为二维字符串。"""
    header_len = len(table.columns)
    while len(table.rows) > start:
        table._tbl.remove(table.rows[-1]._tr)
    if not rows:
        if start == 1 and len(table.rows) == 1:
            blank = table.add_row()
            for c in blank.cells:
                _set_cell(c, "")
        return
    for i, row in enumerate(rows):
        if start + i < len(table.rows):
            cells = table.rows[start + i].cells
        else:
            cells = table.add_row().cells
        for j in range(header_len):
            _set_cell(cells[j], row[j] if j < len(row) else "")


def _replace_paragraphs(doc, mapping):
    for p in doc.paragraphs:
        src = p.text
        if not src:
            continue
        new = src
        for k, v in mapping.items():
            if k in new:
                new = new.replace(k, v)
        if new != src and p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""


def build_docx(agg, title, generator, generated_at) -> bytes:
    from docx import Document

    if not DOCX_TEMPLATE.exists():
        raise HTTPException(status_code=500, detail="报告模板不存在")
    doc = Document(str(DOCX_TEMPLATE))
    p, v = agg["project"], agg["version"]
    date_range = f"{v.start_date or '-'}　～　{v.launch_date or generated_at[:10]}"
    testers = agg["tester_names"] or generator

    if doc.paragraphs:
        _set_run = doc.paragraphs[0]
        if _set_run.runs:
            _set_run.runs[0].text = f"{p.name}{v.version_no}测试报告"
            for r in _set_run.runs[1:]:
                r.text = ""
        else:
            _set_run.text = f"{p.name}{v.version_no}测试报告"
    if len(doc.paragraphs) > 1:
        _set_sub = doc.paragraphs[1]
        if _set_sub.runs:
            _set_sub.runs[0].text = f"（第{agg['flow_round']}轮）"
            for r in _set_sub.runs[1:]:
                r.text = ""
        else:
            _set_sub.text = f"（第{agg['flow_round']}轮）"

    tables = doc.tables
    # 0 文档信息
    if len(tables) > 0:
        info = [
            ["项目名称", p.name],
            ["系统名称", title],
            ["模块名", "、".join(m.name for m in agg["modules"][:8]) or "-"],
            ["测试版本", f"{v.version_no} {v.name or ''}".strip()],
            ["测试周期", date_range],
            ["测试人员", testers],
            ["编写人", generator],
            ["审核人", ""],
            ["批准人", ""],
            ["编写日期", generated_at[:10]],
        ]
        for i, (k, val) in enumerate(info):
            if i < len(tables[0].rows) and len(tables[0].rows[i].cells) >= 2:
                _set_cell(tables[0].rows[i].cells[0], k)
                _set_cell(tables[0].rows[i].cells[1], val)

    # 1 修订记录
    if len(tables) > 1:
        _fill_table(tables[1], [["V1.0", generated_at[:10], generator, "根据系统数据首次生成"]])

    # 2 测试范围（模块）
    if len(tables) > 2:
        rows = []
        for m in agg["modules"] or []:
            rows.append([m.name, m.remark or m.name, "本轮", "-"])
        if not rows:
            rows = [["全部", "本版本全部需求/用例", "本轮", "-"]]
        _fill_table(tables[2], rows)

    # 3 参考文档
    if len(tables) > 3:
        _fill_table(tables[3], [
            ["需求列表", f"{agg['req_total']} 条", "系统"],
            ["测试用例", f"{agg['case_total']} 条", "系统"],
            ["BUG 列表", f"{agg['bug_total']} 条", "系统"],
        ])

    # 4 测试环境（保留模板骨架，只改版本相关）
    if len(tables) > 4 and len(tables[4].rows) > 2:
        _set_cell(tables[4].rows[2].cells[1], f"{p.name} / {v.version_no}")

    # 5 测试进度
    if len(tables) > 5:
        flow = agg["flow"]
        start = flow.started_at.strftime("%Y-%m-%d") if flow and flow.started_at else (v.start_date or "-")
        end = flow.finished_at.strftime("%Y-%m-%d") if flow and flow.finished_at else "-"
        _fill_table(tables[5], [[
            f"第{agg['flow_round']}轮", agg["flow_stage"], start, end, testers, agg["flow_status"],
        ]])

    # 6 用例统计
    if len(tables) > 6:
        er = agg["exec_result"]
        _fill_table(tables[6], [
            ["用例总数", agg["case_total"], "本版本全部用例"],
            ["已执行用例", agg["executed_total"], "有执行记录"],
            ["未执行用例", agg["not_executed"], "尚无执行记录"],
            ["通过用例", agg["pass_count"], "最近一次执行为通过"],
            ["失败用例", agg["fail_count"], "最近一次执行为失败"],
            ["阻塞用例", er.get("阻塞", 0), ""],
            ["用例执行率", _pct(agg["executed_total"], agg["case_total"]), "已执行 / 总数"],
            ["用例通过率", agg["pass_rate"], "通过 / 已执行"],
        ])

    # 7 用例模块分布
    if len(tables) > 7:
        rows, tot, pas, fal = [], 0, 0, 0
        for name, d in agg["by_mod_case"].items():
            tot += d["total"]; pas += d["pass"]; fal += d["fail"]
            rows.append([name, d["total"], d["pass"], d["fail"], _pct(d["pass"], d["pass"] + d["fail"])])
        rows.append(["合计", tot, pas, fal, _pct(pas, pas + fal)])
        _fill_table(tables[7], rows)

    # 8 BUG 总览
    if len(tables) > 8:
        _fill_table(tables[8], [
            ["缺陷总数", agg["bug_total"], "本版本全部 BUG"],
            ["已关闭", agg["bug_closed"], "测试验证通过关闭 / 不是BUG"],
            ["验证通过", agg["bug_status"].get("测试验证通过关闭", 0), ""],
            ["未关闭", agg["bug_open"], "仍在处理中"],
            ["遗留", agg["bug_open"], "上线前需跟踪"],
        ])

    # 9 严重等级
    if len(tables) > 9:
        total = agg["bug_total"] or 1
        rows = []
        for s in SEV_ORDER:
            n = agg["bug_severity"].get(s, 0)
            rows.append([s, n, _pct(n, total)])
        _fill_table(tables[9], rows)

    # 10 发现阶段（对应模板「缺陷来源」）
    if len(tables) > 10:
        total = agg["bug_total"] or 1
        rows = [[k, n, _pct(n, total)] for k, n in agg["bug_stage"].items()] or [["-", 0, "0.0%"]]
        _fill_table(tables[10], rows)

    # 11 模块分布
    if len(tables) > 11:
        rows, tot = [], 0
        for name, n in agg["bug_module"].items():
            closed_n = sum(1 for b in agg["closed"] if (b.module.name if b.module else "未分类") == name)
            rows.append([name, n, closed_n, n - closed_n])
            tot += n
        rows.append(["合计", tot, len(agg["closed"]), agg["bug_open"]])
        _fill_table(tables[11], rows)

    # 12 BUG 明细
    if len(tables) > 12:
        rows = []
        for b in agg["bugs"]:
            rows.append([
                b.bug_no or str(b.id),
                b.title or "",
                _strip_html(b.steps or "") or (b.remark or ""),
                b.found_stage or "",
                b.severity or "",
                b.module.name if b.module else "未分类",
                b.status or "",
            ])
        _fill_table(tables[12], rows or [["-", "无", "", "", "", "", ""]])

    # 13 遗留
    if len(tables) > 13:
        rows = []
        for b in agg["open_bugs"]:
            rows.append([
                b.bug_no or str(b.id),
                b.title or "",
                b.severity or "",
                "跟踪修复",
                b.assignee or b.fixer or "",
                v.launch_date or "",
            ])
        _fill_table(tables[13], rows or [["-", "无遗留", "", "", "", ""]])

    # 14 等级定义保留模板原文

    conclusion = (
        f"本版本共 {agg['case_total']} 条用例，已执行 {agg['executed_total']} 条，通过率 {agg['pass_rate']}；"
        f"BUG {agg['bug_total']} 条，未关闭 {agg['bug_open']} 条。当前流程：{agg['flow_stage']}。"
    )
    advice = "建议关闭全部致命/严重缺陷后再上线。" if agg["bug_open"] else "无未关闭缺陷，可按计划上线。"
    _replace_paragraphs(doc, {
        "测试总结：": f"测试总结：{conclusion}",
        "上线建议：": f"上线建议：{advice}",
    })

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.get("")
def list_reports(page: int = 1, size: int = 10, project_id: int = 0, version_id: int = 0,
                 db: Session = Depends(get_db), _=Depends(get_current_user)):
    q = db.query(TestReport)
    if project_id:
        q = q.filter(TestReport.project_id == project_id)
    if version_id:
        q = q.filter(TestReport.version_id == version_id)
    total = q.count()
    items = q.order_by(TestReport.id.desc()).offset((page - 1) * size).limit(size).all()
    data = []
    for r in items:
        d = row_to_dict(r)
        d["project_name"] = r.project.name if r.project else ""
        d["version_no"] = r.version.version_no if r.version else ""
        d.pop("content", None)
        d.pop("summary", None)
        data.append(d)
    return {"items": data, "total": total}


@router.post("/generate")
def generate_report(payload: dict, db: Session = Depends(get_db), cur=Depends(get_current_user)):
    project_id = payload.get("project_id")
    version_id = payload.get("version_id")
    module_ids = payload.get("module_ids") or []
    agg = _aggregate(db, project_id, version_id, module_ids)
    title = payload.get("title") or f"{agg['project'].name}-{agg['version'].version_no}测试报告"
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    generator = cur.real_name or cur.username
    content = _md_content(title, agg, generator, generated_at)
    r = TestReport(
        title=title, project_id=project_id, version_id=version_id,
        content=content, method="系统", generator=generator,
        module_ids=json.dumps([int(m) for m in module_ids if m]),
    )
    db.add(r)
    db.flush()
    r.report_no = gen_code("RPT", r.id)
    db.commit()
    d = row_to_dict(r)
    d["project_name"] = r.project.name if r.project else ""
    d["version_no"] = r.version.version_no if r.version else ""
    return d


@router.get("/{rid}")
def get_report(rid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    r = db.query(TestReport).filter(TestReport.id == rid).first()
    if not r:
        raise HTTPException(status_code=404, detail="报告不存在")
    d = row_to_dict(r)
    d["project_name"] = r.project.name if r.project else ""
    d["version_no"] = r.version.version_no if r.version else ""
    return d


@router.delete("/{rid}")
def delete_report(rid: int, db: Session = Depends(get_db), _=Depends(require_admin)):
    r = db.query(TestReport).filter(TestReport.id == rid).first()
    if r:
        db.delete(r)
        db.commit()
    return {"ok": True}


@router.get("/{rid}/export")
def export_report(rid: int, db: Session = Depends(get_db), _=Depends(get_current_user)):
    r = db.query(TestReport).filter(TestReport.id == rid).first()
    if not r:
        raise HTTPException(status_code=404, detail="报告不存在")
    try:
        scope_module_ids = json.loads(r.module_ids or "[]")
    except Exception:
        scope_module_ids = []
    agg = _aggregate(db, r.project_id, r.version_id, scope_module_ids)
    generated_at = r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    doc_bytes = build_docx(agg, r.title, r.generator or "", generated_at)
    filename = f"{r.title or r.report_no}.docx"
    encoded = quote(filename)
    return StreamingResponse(
        io.BytesIO(doc_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )
